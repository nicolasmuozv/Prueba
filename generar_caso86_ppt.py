#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Caso 86 — Catálogo de Esquemas Tributarios SII 2025
PPT — Continuación slides 3 en adelante, replicando estilo compañera
Grupo 4: Andrea Añasco, Gema Sepúlveda, Karen Rebolledo, Nicolás Muñoz
Fuentes: DL 824 (LIR), DL 825 (IVA), Código Tributario,
         Circ. N°21/1991, Circ. N°50/2020, Circ. N°65/2015 del SII.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

OUTPUT = "/home/user/Prueba/Caso86_Presentacion.pptx"

# ── Paleta calcada del template de la compañera ──────────────────────────────
FONDO       = RGBColor(0xCB, 0xCB, 0xCB)   # gris claro fondo
CAJA_AZUL   = RGBColor(0x8E, 0xB4, 0xE3)   # azul claro participantes
CAJA_DARK   = RGBColor(0x2B, 0x2B, 0x2B)   # negro/grafito narrativa
CAJA_NARAN  = RGBColor(0xFD, 0xE9, 0xD9)   # naranja claro RESULTADO
ACCENT_NAR  = RGBColor(0xC5, 0x50, 0x1A)   # naranja oscuro título naranja
AMARILLO    = RGBColor(0xE6, 0xA4, 0x3C)   # dorado para headers como compañera
TEXTO_DARK  = RGBColor(0x1C, 0x1C, 0x1C)
TEXTO_CLARO = RGBColor(0xFF, 0xFF, 0xFF)
AZUL_TIT    = RGBColor(0x1A, 0x3A, 0x5C)
GRIS_TXT    = RGBColor(0x55, 0x55, 0x55)
ROJO        = RGBColor(0xB0, 0x30, 0x30)
VERDE       = RGBColor(0x1A, 0x6B, 0x3C)
BLANCO      = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

SW = prs.slide_width
SH = prs.slide_height
BLANK = prs.slide_layouts[6]


# ── Helpers ──────────────────────────────────────────────────────────────────

def set_bg(slide, color=FONDO):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    sp = bg._element.getparent()
    sp.remove(bg._element)
    sp.insert(2, bg._element)
    return bg

def add_title(slide, text, left=Inches(0.6), top=Inches(0.3),
              width=Inches(12.1), height=Inches(0.9),
              size=36, color=TEXTO_DARK, bold=False, italic=False,
              font="Cambria", align=PP_ALIGN.LEFT):
    """Título estilo compañera: serif, grande, gris oscuro, izquierda."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return tb

def add_section_header(slide, text, left, top, width, height=Inches(0.55)):
    """Header rectangular negro con texto naranja subrayado (como 'Normativa Base')."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = CAJA_DARK
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.underline = True
    r.font.color.rgb = AMARILLO
    r.font.name = "Calibri"
    return box

def add_box(slide, text, left, top, width, height,
            fill=CAJA_DARK, font_color=TEXTO_CLARO, size=11,
            bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            shape=MSO_SHAPE.RECTANGLE, no_fill=False, font="Calibri",
            no_line=True):
    box = slide.shapes.add_shape(shape, left, top, width, height)
    if no_fill:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if no_line:
        box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)
    tf.vertical_anchor = valign

    for i, linea in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = linea
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = font_color
        r.font.name = font
    return box

def add_resultado_strip(slide, text, left=Inches(11.3), top=Inches(1.0),
                       width=Inches(1.85), height=Inches(5.7),
                       titulo="RESULTADO"):
    """Banda lateral naranja claro con título naranja oscuro (como compañera)."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = CAJA_NARAN
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = titulo
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = ACCENT_NAR
    r.font.name = "Calibri"

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(8)
    r2 = p2.add_run()
    r2.text = "\n" + text
    r2.font.size = Pt(11)
    r2.font.color.rgb = TEXTO_DARK
    r2.font.name = "Calibri"
    return box


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — ¿Califica la sociedad D? Análisis de requisitos
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Análisis jurídico — Requisitos de la sociedad")

add_section_header(s, "Verificación de los requisitos (Circ. N°21/1991 y N°50/2020 SII)",
                   Inches(0.6), Inches(1.25), Inches(10.5))

# Encabezado tabla
y_hdr = Inches(2.0)
add_box(s, "Requisito copulativo", Inches(0.6), y_hdr, Inches(6.6), Inches(0.5),
        fill=CAJA_DARK, font_color=BLANCO, size=12, bold=True,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)
add_box(s, "¿Se cumple en el Caso 86?", Inches(7.3), y_hdr, Inches(3.85), Inches(0.5),
        fill=CAJA_DARK, font_color=BLANCO, size=12, bold=True,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

requisitos = [
    ("Debe tratarse de una sociedad de personas.",                                   "Cumple", VERDE),
    ("Objeto exclusivo: prestación de servicios o asesorías profesionales.",         "Cumple", VERDE),
    ("Servicios prestados por intermedio de sus socios.",                            "Solo A presta servicios", AMARILLO),
    ("TODOS los socios deben ejercer su profesión para la sociedad.",                "NO CUMPLE — B y C no prestan", ROJO),
    ("No es aceptable que un socio solo aporte capital sin ejercer profesión.",      "NO CUMPLE — B y C solo aportan capital", ROJO),
    ("La realización efectiva de labores profesionales es lo relevante, no la forma de remuneración.", "Solo A realiza labor profesional", ROJO),
]
y = Inches(2.55)
h = Inches(0.6)
for req, est, color in requisitos:
    add_box(s, req, Inches(0.6), y, Inches(6.6), h,
            fill=RGBColor(0xF4, 0xF6, 0xF9), font_color=TEXTO_DARK,
            size=10.5, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)
    add_box(s, est, Inches(7.3), y, Inches(3.85), h,
            fill=color, font_color=BLANCO,
            size=10.5, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    y += h + Inches(0.05)

# Banda lateral RESULTADO
add_resultado_strip(s,
    "La sociedad D NO califica como sociedad de profesionales del Art. 42 N°2 inc. 3° LIR, "
    "al incumplir requisitos copulativos de las Circulares N°21/1991 y N°50/2020.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.7))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Normas Generales Antielusivas (NGA)
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Normas Generales Antielusivas (NGA)")

add_section_header(s, "Marco aplicable — Código Tributario",
                   Inches(0.6), Inches(1.25), Inches(10.5))

# Cuatro cards 2x2
art_data = [
    ("Art. 4 bis CT",  "Principio de buena fe",
     "Las obligaciones tributarias nacen y se exigen con arreglo a la naturaleza jurídica "
     "de los hechos, cualquiera sea la forma o denominación que los interesados les hubieren dado."),
    ("Art. 4 ter CT",  "Abuso de las formas jurídicas",
     "Hay abuso cuando se evita total o parcialmente el hecho gravado, o se disminuye la "
     "base imponible u obligación, mediante actos que no produzcan resultados o efectos "
     "jurídicos o económicos relevantes distintos de los meramente tributarios."),
    ("Art. 4 quáter CT","Simulación",
     "Habrá simulación cuando los actos o negocios disimulen la configuración del hecho "
     "gravado o la naturaleza de los elementos constitutivos de la obligación tributaria, "
     "o su verdadero monto o data de nacimiento."),
    ("Art. 4 quinquies CT","Procedimiento",
     "La existencia de abuso o simulación será declarada por el Tribunal Tributario y "
     "Aduanero, a requerimiento del Director del SII. Instrucciones: Circular N°65 de 2015."),
]

x0 = Inches(0.6)
y0 = Inches(2.0)
w_c = Inches(5.25)
h_c = Inches(2.35)
for i, (cod, tit, txt) in enumerate(art_data):
    col = i % 2
    fila = i // 2
    x = x0 + col * (w_c + Inches(0.2))
    y = y0 + fila * (h_c + Inches(0.2))
    # Header negro con texto amarillo (estilo compañera)
    add_section_header(s, f"{cod}  ·  {tit}", x, y, w_c, height=Inches(0.5))
    # Cuerpo
    add_box(s, txt, x, y + Inches(0.5), w_c, h_c - Inches(0.5),
            fill=RGBColor(0xF4, 0xF6, 0xF9), font_color=TEXTO_DARK,
            size=11, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP)

# Banda lateral
add_resultado_strip(s,
    "Las NGA fueron incorporadas por la Ley N°20.780 (2014) y modificadas por la "
    "Ley N°21.210 (2020). La Circular N°65 de 2015 imparte instrucciones para su aplicación.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.9))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — ¿Abuso o Simulación? Calificación jurídica
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "¿Abuso o Simulación?")

add_section_header(s, "Calificación del esquema bajo las NGA",
                   Inches(0.6), Inches(1.25), Inches(10.5))

# Dos columnas comparativas
lx = Inches(0.6)
rx = Inches(5.85)
cw = Inches(5.05)

# ABUSO header
add_section_header(s, "ABUSO  (Art. 4 ter CT)", lx, Inches(2.0), cw, height=Inches(0.5))
add_box(s,
    "• Los actos son REALES: la sociedad D existe, los socios son reales, hay capital "
    "aportado y distribución efectiva de utilidades.\n\n"
    "• La FORMA jurídica utilizada es inapropiada para el sustrato económico: la "
    "sociedad de profesionales no está prevista para canalizar el trabajo individual "
    "de un solo socio.\n\n"
    "• B y C no producen efectos jurídicos ni económicos relevantes distintos del "
    "meramente tributario.\n\n"
    "• Único efecto: reducción de la progresión del IGC del socio A.",
    lx, Inches(2.5), cw, Inches(4.2),
    fill=RGBColor(0xFB, 0xE5, 0xE5), font_color=TEXTO_DARK,
    size=11, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP)

# SIMULACIÓN header
add_section_header(s, "SIMULACIÓN  (Art. 4 quáter CT)", rx, Inches(2.0), cw, height=Inches(0.5))
add_box(s,
    "• Requiere divergencia consciente entre la voluntad real y la declarada.\n\n"
    "• Supone la existencia de un acto secreto u oculto que se disimula bajo otro "
    "aparente (negocio jurídico simulado).\n\n"
    "• En el Caso 86 NO existe acto disimulado: la sociedad y los socios son lo que "
    "aparentan ser; los actos societarios y la distribución de utilidades son reales.\n\n"
    "• Se descarta la figura de simulación para este esquema.",
    rx, Inches(2.5), cw, Inches(4.2),
    fill=RGBColor(0xFD, 0xF0, 0xE6), font_color=TEXTO_DARK,
    size=11, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP)

# Banda RESULTADO
add_resultado_strip(s,
    "Posición del grupo:\n\nEl Caso 86 configura ABUSO de las formas jurídicas (Art. 4 ter CT), "
    "no simulación.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.7))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Fundamentos del abuso (los 3 elementos del Art. 4 ter)
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Fundamentos del abuso — Art. 4 ter CT")

add_section_header(s, "Los tres elementos están presentes en el Caso 86",
                   Inches(0.6), Inches(1.25), Inches(10.5))

elementos = [
    ("Elemento 1 — Evitar o disminuir la obligación tributaria",
     "El esquema diluye los ingresos del socio A en tres tramos independientes de IGC, "
     "reduciendo la tasa marginal efectiva. La carga tributaria del contribuyente A "
     "disminuye en relación a la que correspondería si los ingresos los percibiera "
     "directamente como independiente del Art. 42 N°2 LIR o como único prestador del servicio."),
    ("Elemento 2 — Ausencia de efectos jurídicos o económicos relevantes",
     "La participación de B y C no produce efectos relevantes distintos del tributario:\n"
     "•  No prestan servicios profesionales para la sociedad.\n"
     "•  No tienen rol estratégico ni decisional.\n"
     "•  El capital aportado no es significativo en relación a los ingresos.\n"
     "•  Su única función es absorber utilidades para diluir el IGC de A."),
    ("Elemento 3 — Forma jurídica inapropiada para el sustrato económico",
     "La figura 'sociedad de profesionales' (Art. 42 N°2 inc. 3° LIR) está prevista para "
     "profesionales que efectivamente trabajan en conjunto, según lo precisan las Circulares "
     "N°21/1991 y N°50/2020 del SII. Utilizarla para canalizar el trabajo individual de un "
     "solo socio constituye una utilización abusiva de la institución jurídica."),
]
y_e = Inches(2.0)
for tit, txt in elementos:
    h_e = Inches(1.55)
    add_section_header(s, tit, Inches(0.6), y_e, Inches(10.5), height=Inches(0.5))
    add_box(s, txt, Inches(0.6), y_e + Inches(0.5), Inches(10.5), h_e - Inches(0.5),
            fill=RGBColor(0xF4, 0xF6, 0xF9), font_color=TEXTO_DARK,
            size=10.5, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP)
    y_e += h_e + Inches(0.08)

add_resultado_strip(s,
    "Configurándose los tres elementos del Art. 4 ter CT, el Director del SII puede "
    "requerir al TTA la declaración de abuso (Art. 4 quinquies CT).",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.95))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Efecto tributario (ilustrativo)
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Efecto tributario del esquema")

add_section_header(s, "Comparación ilustrativa — Dilución de la progresión del IGC",
                   Inches(0.6), Inches(1.25), Inches(10.5))

# Encabezado tabla
y_h = Inches(2.0)
cols_t = [
    ("Escenario",                  Inches(0.6),  Inches(3.0)),
    ("Atribución de la renta",     Inches(3.65), Inches(3.7)),
    ("Tramo marginal IGC aprox.",  Inches(7.4),  Inches(2.0)),
    ("Resultado",                  Inches(9.45), Inches(1.65)),
]
for txt, l, w in cols_t:
    add_box(s, txt, l, y_h, w, Inches(0.5),
            fill=CAJA_DARK, font_color=BLANCO, size=11, bold=True,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

filas_t = [
    ("Sin esquema\n(A trabaja solo, sin sociedad)",
     "A declara el total de los ingresos",
     "Tramo alto (≈40%)",
     "Tributación\ncorrecta",
     RGBColor(0xE8, 0xF3, 0xE8)),
    ("Con esquema D\n(3 socios, capital simbólico)",
     "A : B : C en partes iguales (33,33% c/u)",
     "Cada socio en tramo medio",
     "Ahorro fiscal\nartificial",
     RGBColor(0xFB, 0xE5, 0xE5)),
    ("Diferencial",
     "Los ingresos los genera íntegramente el socio A",
     "Disminución total del IGC",
     "ABUSO\nArt. 4 ter CT",
     RGBColor(0xFD, 0xE9, 0xD9)),
]
y_t = Inches(2.55)
for esc, atrib, tramo, res, color in filas_t:
    h = Inches(1.1)
    add_box(s, esc,   Inches(0.6),  y_t, Inches(3.0),  h,
            fill=color, font_color=TEXTO_DARK, size=10.5, bold=True,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)
    add_box(s, atrib, Inches(3.65), y_t, Inches(3.7),  h,
            fill=color, font_color=TEXTO_DARK, size=10.5,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)
    add_box(s, tramo, Inches(7.4),  y_t, Inches(2.0),  h,
            fill=color, font_color=TEXTO_DARK, size=10.5,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_box(s, res,   Inches(9.45), y_t, Inches(1.65), h,
            fill=color, font_color=TEXTO_DARK, size=10, bold=True,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    y_t += h + Inches(0.05)

# Nota
add_box(s, "Nota: las tasas son ilustrativas. El perjuicio fiscal real dependerá del "
        "monto efectivo de los ingresos y otras rentas de cada socio.",
        Inches(0.6), Inches(6.05), Inches(10.5), Inches(0.4),
        no_fill=True, font_color=GRIS_TXT, size=10)

add_resultado_strip(s,
    "Declarado el abuso, el efecto es la recalificación: atribución íntegra de los "
    "ingresos al socio A para el cálculo del IGC, con intereses y multas.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.95))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Alternativas de solución (parte 1)
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Alternativas legítimas de solución (1/2)")

add_section_header(s, "Estructuras conforme al ordenamiento tributario",
                   Inches(0.6), Inches(1.25), Inches(10.5))

alts_1 = [
    ("Alternativa 1 — Profesional independiente (Art. 42 N°2 LIR)",
     "El socio A ejerce como persona natural independiente, emite boletas de honorarios "
     "y declara en segunda categoría. Puede optar entre gastos efectivos o presuntos "
     "(30% con tope de 15 UTA). Tributa con IGC sobre la totalidad de los ingresos. "
     "Opción simple, sin exposición a NGA."),
    ("Alternativa 2 — Sociedad de profesionales con socios que ejerzan profesión",
     "Si A desea asociarse, debe hacerlo con profesionales que efectivamente coadyuven "
     "a la prestación de servicios para la sociedad. En ese caso D califica legítimamente "
     "como sociedad de profesionales y puede optar por tributar en primera o segunda "
     "categoría conforme al Art. 42 N°2 inc. 3° LIR."),
    ("Alternativa 3 — Distribución de utilidades proporcional al trabajo real",
     "Manteniendo a A, B y C como socios, los estatutos pueden contemplar:\n"
     "(i) participación social diferenciada acorde al aporte efectivo (ej. A 90%, B 5%, C 5%); o\n"
     "(ii) sueldo patronal a A previo al reparto de utilidades.\n"
     "La distribución debe reflejar la realidad económica del aporte."),
]
y_a = Inches(2.0)
for tit, txt in alts_1:
    h_a = Inches(1.55)
    add_section_header(s, tit, Inches(0.6), y_a, Inches(10.5), height=Inches(0.5))
    add_box(s, txt, Inches(0.6), y_a + Inches(0.5), Inches(10.5), h_a - Inches(0.5),
            fill=RGBColor(0xF4, 0xF6, 0xF9), font_color=TEXTO_DARK,
            size=10.5, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP)
    y_a += h_a + Inches(0.08)

add_resultado_strip(s,
    "Toda alternativa debe poder acreditar SUSTANCIA ECONÓMICA REAL. La sola "
    "existencia formal de la figura no es suficiente.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.95))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — Alternativas de solución (parte 2)
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Alternativas legítimas de solución (2/2)")

add_section_header(s, "Vehículos societarios con sustancia económica",
                   Inches(0.6), Inches(1.25), Inches(10.5))

alts_2 = [
    ("Alternativa 4 — SpA acogida al régimen Pro Pyme 14 D N°3 LIR",
     "Constitución de Sociedad por Acciones (SpA) acogida al régimen Pro Pyme General "
     "del Art. 14 letra D N°3 de la LIR, con tasa de IDPC del 25%. Permite planificación "
     "de retiros y reinversión. No tiene las restricciones específicas de la sociedad de "
     "profesionales del Art. 42 N°2, pero exige sustancia económica real en la sociedad."),
    ("Alternativa 5 — Régimen Pro Pyme Transparente — 14 D N°8 LIR",
     "Si los socios son únicamente personas naturales, puede evaluarse el régimen "
     "transparente del Art. 14 letra D N°8 LIR: la empresa no paga IDPC y los socios "
     "tributan directamente con IGC. ATENCIÓN: este régimen no resuelve el problema "
     "de fondo si B y C siguen siendo socios sin sustancia económica real; el SII "
     "puede igualmente aplicar las NGA."),
]
y_a = Inches(2.0)
for tit, txt in alts_2:
    h_a = Inches(2.0)
    add_section_header(s, tit, Inches(0.6), y_a, Inches(10.5), height=Inches(0.5))
    add_box(s, txt, Inches(0.6), y_a + Inches(0.5), Inches(10.5), h_a - Inches(0.5),
            fill=RGBColor(0xF4, 0xF6, 0xF9), font_color=TEXTO_DARK,
            size=10.5, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP)
    y_a += h_a + Inches(0.1)

add_resultado_strip(s,
    "La elección de la estructura dependerá de los objetivos del contribuyente, "
    "el nivel de ingresos y la composición real del equipo profesional.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.95))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — Conclusiones
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_title(s, "Conclusiones")

add_section_header(s, "Caso 86 — Síntesis del análisis del grupo",
                   Inches(0.6), Inches(1.25), Inches(10.5))

conclusiones = [
    "1.  El esquema configura ABUSO de las formas jurídicas (Art. 4 ter CT). No es "
    "simulación: los actos son reales, pero la forma jurídica es inapropiada para el "
    "sustrato económico que se quiere amparar.",
    "2.  La sociedad D NO califica como sociedad de profesionales conforme a las "
    "Circulares N°21/1991 y N°50/2020 del SII, dado que B y C no ejercen su profesión "
    "para la sociedad.",
    "3.  Configurándose los tres elementos del Art. 4 ter CT, el Director del SII puede "
    "requerir al Tribunal Tributario y Aduanero la declaración de abuso (Art. 4 quinquies CT), "
    "con atribución íntegra de los ingresos al socio A, más intereses y multas.",
    "4.  Existen alternativas legítimas de organización tributaria (profesional independiente, "
    "sociedad de profesionales con socios reales, distribución proporcional al trabajo, "
    "SpA Pro Pyme).",
    "5.  Toda estructura debe poder acreditar SUSTANCIA ECONÓMICA REAL. La sola forma "
    "jurídica no es suficiente; el SII y los tribunales exigirán demostración del aporte "
    "efectivo de cada socio (trabajo, capital, clientela o dirección).",
]
y_c = Inches(2.0)
for txt in conclusiones:
    h_c = Inches(0.85)
    add_box(s, txt, Inches(0.6), y_c, Inches(10.5), h_c,
            fill=RGBColor(0xF4, 0xF6, 0xF9), font_color=TEXTO_DARK,
            size=11, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE)
    y_c += h_c + Inches(0.07)

add_resultado_strip(s,
    "El Caso 86 ilustra el límite entre la planificación tributaria legítima y el "
    "abuso: la forma jurídica debe estar respaldada por sustancia económica real.",
    left=Inches(11.3), top=Inches(2.0), width=Inches(1.85), height=Inches(4.95))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — Cierre / Grupo 4
# ═══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
set_bg(s, CAJA_DARK)

add_title(s, "Grupo 4", left=Inches(1), top=Inches(1.5), width=Inches(11.3),
          height=Inches(1.0), size=48, color=AMARILLO, bold=True,
          font="Cambria", align=PP_ALIGN.CENTER)

add_title(s, "Caso 86 — Catálogo de Esquemas Tributarios SII 2025",
          left=Inches(1), top=Inches(2.7), width=Inches(11.3), height=Inches(0.6),
          size=18, color=BLANCO, font="Calibri", align=PP_ALIGN.CENTER)

# Integrantes
miembros = ["Andrea Añasco", "Gema Sepúlveda", "Karen Rebolledo", "Nicolás Muñoz"]
x0 = Inches(1.5)
w  = Inches(2.45)
for i, m in enumerate(miembros):
    add_box(s, m, x0 + i * (w + Inches(0.18)), Inches(4.0), w, Inches(0.7),
            fill=CAJA_NARAN, font_color=ACCENT_NAR,
            size=14, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

add_title(s, "Magíster en Dirección Tributaria — Universidad Viña del Mar",
          left=Inches(1), top=Inches(6.3), width=Inches(11.3), height=Inches(0.5),
          size=12, color=RGBColor(0xCC, 0xCC, 0xCC), font="Calibri", align=PP_ALIGN.CENTER)


prs.save(OUTPUT)
print(f"✔ Presentación PPT generada: {OUTPUT}")
