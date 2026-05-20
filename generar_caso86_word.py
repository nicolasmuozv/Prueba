#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Caso 86 — Catálogo de Esquemas Tributarios SII 2025
Informe escrito v2 — Enfocado en el caso, con ejemplos y tablas
Magíster en Dirección Tributaria UVM — Grupo 4
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "/home/user/Prueba/Caso86_Informe.docx"

AZUL    = RGBColor(0x1A, 0x3A, 0x5C)
CELESTE = RGBColor(0x2E, 0x6D, 0xA4)
ROJO    = RGBColor(0xB0, 0x30, 0x30)
GRIS    = RGBColor(0x55, 0x55, 0x55)
NEGRO   = RGBColor(0x1C, 0x1C, 0x1C)
VERDE   = RGBColor(0x1A, 0x6B, 0x3C)
NARANJA = RGBColor(0xC5, 0x50, 0x1A)

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.paragraph_format.space_after   = Pt(4)
style.paragraph_format.line_spacing  = 1.15


# ── Helpers ──────────────────────────────────────────────────────────────────

def set_cell_bg(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  color_hex)
    tcPr.append(shd)

def add_title(text, size=16, color=AZUL, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.font.size  = Pt(size)
    r.font.bold  = True
    r.font.color.rgb = color
    return p

def add_section(text, size=12, color=AZUL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    r.font.size  = Pt(size)
    r.font.bold  = True
    r.font.color.rgb = color
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '8')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '2E6DA4')
    pBdr.append(bot)
    pPr.append(pBdr)
    return p

def add_subtitle(text, size=11, color=CELESTE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(text)
    r.font.size  = Pt(size)
    r.font.bold  = True
    r.font.color.rgb = color
    return p

def add_para(text, bold=False, justify=True, color=None, size=10.5, italic=False):
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.size   = Pt(size)
    r.font.bold   = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = color
    return p

def add_bullet(text, level=0, color=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent  = Cm(0.5 + 0.5 * level)
    p.paragraph_format.space_after  = Pt(2)
    run = p.runs[0] if p.runs else p.add_run("")
    run.text           = text
    run.font.size      = Pt(10)
    if color:
        run.font.color.rgb = color
    return p

def add_box_para(text, bg_hex="EBF3FB", bold=False, color=AZUL, size=10.5):
    """Párrafo con fondo de color (simulado con tabla de 1 celda)."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    set_cell_bg(cell, bg_hex)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run(text)
    r.font.size  = Pt(size)
    r.font.bold  = bold
    r.font.color.rgb = color
    doc.add_paragraph()
    return t

def row_hdr(table, texts, bg="1A3A5C"):
    row = table.rows[0].cells
    for i, txt in enumerate(texts):
        row[i].text = txt
        set_cell_bg(row[i], bg)
        for p in row[i].paragraphs:
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.bold      = True
                r.font.size      = Pt(10)

def row_data(table, row_idx, vals, bg=None, bold=False, size=9.5):
    row = table.rows[row_idx].cells
    for i, val in enumerate(vals):
        row[i].text = val
        if bg:
            set_cell_bg(row[i], bg)
        for p in row[i].paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.bold = bold


# ═══════════════════════════════════════════════════════════════════════════════
# PORTADA
# ═══════════════════════════════════════════════════════════════════════════════

enc = doc.add_paragraph()
enc.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_enc = enc.add_run("MAGÍSTER EN DIRECCIÓN TRIBUTARIA — UNIVERSIDAD VIÑA DEL MAR")
r_enc.font.size  = Pt(9)
r_enc.font.bold  = True
r_enc.font.color.rgb = GRIS

add_title("CASO 86", size=20, color=AZUL)
add_title("Prestación de Servicios Profesionales mediante Sociedad Interpuesta", size=12, color=CELESTE)
add_title("Catálogo de Esquemas Tributarios SII 2025", size=10, color=GRIS)

doc.add_paragraph()

# Tabla grupo
t_grupo = doc.add_table(rows=2, cols=2)
t_grupo.style = 'Light Grid Accent 1'
t_grupo.alignment = WD_TABLE_ALIGNMENT.CENTER

set_cell_bg(t_grupo.rows[0].cells[0], "1A3A5C")
set_cell_bg(t_grupo.rows[0].cells[1], "1A3A5C")
t_grupo.rows[0].cells[0].text = "Grupo 4"
t_grupo.rows[0].cells[1].text = "Magíster en Dirección Tributaria UVM"
for c in t_grupo.rows[0].cells:
    for p in c.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.bold = True
            r.font.size = Pt(11)

integrantes = "Andrea Añasco  ·  Gema Sepúlveda  ·  Karen Rebolledo  ·  Nicolás Muñoz"
t_grupo.rows[1].cells[0].merge(t_grupo.rows[1].cells[1])
t_grupo.rows[1].cells[0].text = integrantes
for p in t_grupo.rows[1].cells[0].paragraphs:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = AZUL

doc.add_paragraph()

linea = doc.add_paragraph()
linea.alignment = WD_ALIGN_PARAGRAPH.CENTER
linea.add_run("─" * 80).font.color.rgb = CELESTE


# ═══════════════════════════════════════════════════════════════════════════════
# I. EL ESQUEMA — ¿QUÉ SE HACE Y POR QUÉ IMPORTA?
# ═══════════════════════════════════════════════════════════════════════════════
add_section("I.  EL ESQUEMA — ¿QUÉ SE HACE Y POR QUÉ IMPORTA?")

add_para(
    "El Caso 86 describe una práctica frecuente en el mundo de los profesionales liberales: "
    "constituir una sociedad de profesionales con socios que, en la realidad, nunca trabajan "
    "para la sociedad. El objetivo concreto es fraccionar el ingreso del profesional activo "
    "en varios RUT y, con ello, bajar artificialmente el tramo del Impuesto Global Complementario (IGC)."
)

add_subtitle("El caso concreto")

add_para(
    "Pedro, médico cirujano, genera honorarios anuales por $300 millones. Si los declara "
    "directamente bajo el Art. 42 N°2 LIR, paga IGC en el tramo más alto (40%). "
    "Para reducir esa carga, constituye junto a su señora y su cuñado —que son profesores "
    "y nunca atienden pacientes— la 'Sociedad Profesional Médica Ltda.', con participación "
    "del 33,33% cada uno. Los ingresos de Pedro entran a la sociedad y se reparten en partes "
    "iguales: cada uno declara $100 millones y queda en un tramo de IGC más bajo."
)

# Tabla esquema visual
add_subtitle("El flujo del esquema")

t_flujo = doc.add_table(rows=4, cols=3)
t_flujo.style = 'Table Grid'
t_flujo.alignment = WD_TABLE_ALIGNMENT.CENTER

row_hdr(t_flujo, ["Paso", "¿Qué ocurre?", "Sujeto"])
filas_flujo = [
    ["1", "Pedro presta servicios médicos y factura $300 MM", "Socio A (Pedro)"],
    ["2", "Los ingresos ingresan a la Sociedad D", "Sociedad de profesionales"],
    ["3", "Las utilidades se distribuyen en partes iguales ($100 MM c/u)", "A, B y C — todos declaran IGC"],
]
for i, f in enumerate(filas_flujo):
    row_data(t_flujo, i + 1, f, bg="EBF3FB" if i % 2 == 0 else None)

doc.add_paragraph()

add_box_para(
    "RESULTADO: lo que en manos de Pedro tributaría al 40%, queda diluido en tres tramos "
    "independientes. El ahorro fiscal puede superar los $20–40 millones anuales, sin que "
    "B ni C hayan atendido un solo paciente.",
    bg_hex="FDE9D9", color=NARANJA, bold=True
)

add_subtitle("Diagrama simplificado del esquema")

# Diagrama en tabla
t_diag = doc.add_table(rows=3, cols=5)
t_diag.style = 'Table Grid'
t_diag.alignment = WD_TABLE_ALIGNMENT.CENTER

celdas = [
    (0, 0, "Pedro\n(médico)", "2E6DA4"),
    (0, 2, "Sociedad D\n(profesionales)", "1A3A5C"),
    (0, 4, "SII recibe\nIGC de 3 RUT", "2E6DA4"),
    (1, 1, "→ honorarios", None),
    (1, 3, "→ distribuye", None),
    (2, 2, "B (33%) · C (33%) · A (33%)", "555555"),
]
for row_i, col_i, txt, bg in celdas:
    cell = t_diag.rows[row_i].cells[col_i]
    cell.text = txt
    if bg:
        set_cell_bg(cell, bg)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.bold = True
                r.font.size = Pt(9.5)
    else:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.color.rgb = GRIS

doc.add_paragraph()


# ═══════════════════════════════════════════════════════════════════════════════
# II. ¿QUÉ DICE EL DERECHO? — SÍNTESIS NORMATIVA
# ═══════════════════════════════════════════════════════════════════════════════
add_section("II.  ¿QUÉ DICE EL DERECHO? — SÍNTESIS NORMATIVA")

add_para(
    "No se trata de hacer un repaso de la ley, sino de identificar qué normas son relevantes "
    "para calificar el esquema y cuáles le aplican directamente. La tabla siguiente sintetiza "
    "el marco legal pertinente:"
)

t_norma = doc.add_table(rows=6, cols=3)
t_norma.style = 'Light Grid Accent 1'
t_norma.alignment = WD_TABLE_ALIGNMENT.CENTER

row_hdr(t_norma, ["Norma", "Contenido clave", "¿Cómo aplica al Caso 86?"])
normas = [
    ["Art. 42 N°2 LIR",
     "Rentas de segunda categoría. Las sociedades de profesionales pueden optar por 1ª categoría "
     "si TODOS los socios ejercen su profesión para ella.",
     "D no cumple: B y C no ejercen profesión para la sociedad → no califica como sociedad de profesionales."],
    ["Circ. N°21/1991 y N°50/2020 SII",
     "Requisitos copulativos para ser sociedad de profesionales. Todos los socios deben prestar "
     "servicios efectivos; no es válido que uno solo aporte capital.",
     "Incumplimiento flagrante: B y C solo aportan capital."],
    ["Art. 4 bis CT",
     "Las obligaciones tributarias nacen según la naturaleza jurídica real de los actos, "
     "no según la forma que les dan las partes.",
     "Obligaciones de Pedro no se alteran por usar la sociedad como vehículo."],
    ["Art. 4 ter CT",
     "ABUSO: actos que evitan o disminuyen el hecho gravado sin producir efectos "
     "jurídicos/económicos relevantes distintos del tributario.",
     "El esquema reduce la carga de Pedro sin que B y C aporten valor real → ABUSO."],
    ["Art. 4 quinquies CT",
     "El Director del SII puede requerir al Tribunal Tributario y Aduanero (TTA) "
     "la declaración de abuso.",
     "Consecuencia práctica: recalificación + atribución íntegra a Pedro + intereses + multas."],
]
for i, f in enumerate(normas):
    row_data(t_norma, i + 1, f, bg="F4F6F9" if i % 2 == 0 else None)

doc.add_paragraph()


# ═══════════════════════════════════════════════════════════════════════════════
# III. CALIFICACIÓN JURÍDICA — ABUSO vs. SIMULACIÓN
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_page_break()
add_section("III.  CALIFICACIÓN JURÍDICA — ¿ABUSO O SIMULACIÓN?")

add_para(
    "Esta es la pregunta central del caso. El Código Tributario distingue dos figuras "
    "antielusivas distintas con efectos procesales diferentes. La correcta calificación "
    "importa porque define qué debe probar el SII y cuál norma aplica."
)

# Tabla comparativa
t_cal = doc.add_table(rows=6, cols=3)
t_cal.style = 'Light Grid Accent 1'
t_cal.alignment = WD_TABLE_ALIGNMENT.CENTER

row_hdr(t_cal, ["Criterio", "ABUSO (Art. 4 ter CT)", "SIMULACIÓN (Art. 4 quáter CT)"])
comparacion = [
    ["¿Los actos son reales?",
     "SÍ — la sociedad existe, hay socios y distribución real",
     "NO — hay un acto declarado que disimula otro real y secreto"],
    ["¿Hay divergencia entre lo declarado y lo real?",
     "NO — el contrato dice exactamente lo que se hace",
     "SÍ — hay voluntad real oculta distinta de la declarada"],
    ["¿Qué se cuestiona?",
     "La FORMA jurídica elegida (inapropiada para el sustrato económico)",
     "El CONTENIDO del acto (simula algo que no existe)"],
    ["¿Qué debe probar el SII?",
     "Ausencia de efectos jurídicos/económicos relevantes distintos del tributario",
     "Existencia de un acto oculto o disimulado"],
    ["¿Aplica al Caso 86?",
     "SÍ → B y C no aportan valor real; la sociedad sirve solo para diluir el IGC de Pedro",
     "NO → la sociedad y los socios son reales; no hay acto secreto que ocultar"],
]
for i, f in enumerate(comparacion):
    bg = "F4F6F9" if i % 2 == 0 else None
    row_data(t_cal, i + 1, f, bg=bg)
    # Marcar columna ABUSO si aplica
    if i == 4:
        set_cell_bg(t_cal.rows[i + 1].cells[1], "E8F5E9")
        set_cell_bg(t_cal.rows[i + 1].cells[2], "FBEAEA")

doc.add_paragraph()

add_box_para(
    "CONCLUSIÓN DEL GRUPO: El Caso 86 es ABUSO de las formas jurídicas (Art. 4 ter CT). "
    "Los actos son reales pero se utiliza una figura societaria de forma inapropiada, "
    "sin sustancia económica real para B y C, con el único propósito de reducir el IGC de Pedro.",
    bg_hex="EBF3FB", color=AZUL, bold=True
)


# ═══════════════════════════════════════════════════════════════════════════════
# IV. ¿POR QUÉ ES ABUSO? — LOS TRES ELEMENTOS
# ═══════════════════════════════════════════════════════════════════════════════
add_section("IV.  ¿POR QUÉ ES ABUSO? — LOS TRES ELEMENTOS DEL ART. 4 TER CT")

add_para("El Art. 4 ter CT exige que concurran tres elementos. En el Caso 86, los tres están presentes:")

add_subtitle("Elemento 1 — Reducción de la obligación tributaria")
add_para(
    "Pedro, con ingresos de $300 millones, debería tributar en el tramo más alto del IGC. "
    "Con el esquema, cada socio declara solo $100 millones, cayendo en un tramo inferior. "
    "El ahorro fiscal no es consecuencia de planificación legítima, sino de fraccionar "
    "artificialmente un ingreso que económicamente pertenece solo a Pedro."
)

add_subtitle("Elemento 2 — Ausencia de efectos jurídicos o económicos relevantes")
add_para("¿Qué aportan realmente B y C a la sociedad D? Ninguna de las siguientes cosas:")

t_aporte = doc.add_table(rows=5, cols=2)
t_aporte.style = 'Light Grid Accent 1'
t_aporte.alignment = WD_TABLE_ALIGNMENT.CENTER

row_hdr(t_aporte, ["¿Aportan B y C...?", "Respuesta"])
aportes = [
    ["Servicios profesionales o atención de clientes",    "✖  No. Solo Pedro trabaja."],
    ["Capital significativo respecto a los ingresos",     "✖  Capital simbólico; los ingresos son $300 MM."],
    ["Clientela, contactos o cartera propia",             "✖  No."],
    ["Dirección estratégica o decisiones empresariales",  "✖  No. Solo reciben retiros de utilidades."],
]
for i, f in enumerate(aportes):
    row_data(t_aporte, i + 1, f, bg="FBEAEA" if i % 2 == 0 else None)

doc.add_paragraph()

add_subtitle("Elemento 3 — Forma jurídica inapropiada para el sustrato económico")
add_para(
    "La sociedad de profesionales es una figura creada para que varios profesionales se "
    "organicen conjuntamente. No está pensada para canalizar el trabajo individual de uno solo "
    "y repartir sus utilidades entre socios pasivos. Usar esa forma jurídica para ese propósito "
    "es exactamente lo que el Art. 4 ter CT define como abuso."
)


# ═══════════════════════════════════════════════════════════════════════════════
# V. EFECTO TRIBUTARIO — EJEMPLO CON NÚMEROS
# ═══════════════════════════════════════════════════════════════════════════════
add_section("V.  EFECTO TRIBUTARIO — EJEMPLO CON NÚMEROS")

add_para(
    "Para ilustrar el perjuicio fiscal, consideremos a Pedro con honorarios anuales de $300 MM "
    "y una tabla simplificada del IGC. Los tramos son aproximados a efectos ilustrativos:"
)

t_ej = doc.add_table(rows=4, cols=4)
t_ej.style = 'Light Grid Accent 1'
t_ej.alignment = WD_TABLE_ALIGNMENT.CENTER

row_hdr(t_ej, ["Escenario", "Renta declarada por Pedro", "Tramo IGC aplicable", "Resultado"])
ejemplos = [
    ["Sin esquema\n(Pedro trabaja solo)", "$300 MM (Pedro declara todo)",
     "Tramo ≈ 40% (tramo máximo)",
     "Pedro paga IGC alto"],
    ["Con esquema D\n(3 socios iguales)", "$100 MM (Pedro · B · C, c/u declara 1/3)",
     "Cada socio en tramo ≈ 23–35%",
     "Ahorro artificial del IGC de Pedro"],
    ["Diferencial estimado", "Pedro genera el 100% de los $300 MM",
     "Diferencia de tramos = ahorro real",
     "PERJUICIO FISCAL → ABUSO"],
]
for i, f in enumerate(ejemplos):
    bg = "EBF3FB" if i == 0 else ("FBEAEA" if i == 1 else "FDE9D9")
    row_data(t_ej, i + 1, f, bg=bg)

doc.add_paragraph()

add_para(
    "Si el SII aplica exitosamente el Art. 4 quinquies CT, el efecto para Pedro es: "
    "(1) se le atribuye el total de $300 MM para su IGC; (2) se aplican intereses y multas "
    "sobre la diferencia de impuestos no pagados; (3) los retiros percibidos por B y C "
    "son recalificados en su propio IGC según corresponda.",
    color=ROJO, bold=True
)


# ═══════════════════════════════════════════════════════════════════════════════
# VI. ALTERNATIVAS LEGÍTIMAS DE SOLUCIÓN
# ═══════════════════════════════════════════════════════════════════════════════
doc.add_page_break()
add_section("VI.  ¿QUÉ PUEDE HACER PEDRO? — ALTERNATIVAS LEGÍTIMAS")

add_para(
    "El problema no es ejercer como sociedad, sino hacerlo de forma que carezca de sustancia "
    "económica real. Existen alternativas legítimas que Pedro puede adoptar sin exponerse a NGA:"
)

t_alt = doc.add_table(rows=6, cols=3)
t_alt.style = 'Light Grid Accent 1'
t_alt.alignment = WD_TABLE_ALIGNMENT.CENTER

row_hdr(t_alt, ["Alternativa", "¿En qué consiste?", "Ventaja / Condición"])
alternativas = [
    ["Alt. 1\nProfesional independiente",
     "Pedro emite boletas de honorarios directamente. Paga IGC sobre sus ingresos. "
     "Puede deducir gastos efectivos o presuntos (30%, tope 15 UTA).",
     "Sencillo, sin riesgo de NGA. Sin beneficio societario."],
    ["Alt. 2\nSociedad con socios que trabajen",
     "Pedro se asocia con otros médicos u otros profesionales que efectivamente "
     "atiendan pacientes o presten servicios para la sociedad.",
     "Califica legítimamente como sociedad de profesionales. "
     "La distribución de utilidades refleja trabajo real."],
    ["Alt. 3\nDistribución proporcional al trabajo",
     "Si D se mantiene con A, B y C, los estatutos establecen que Pedro recibe "
     "el 90% de las utilidades, y B y C un 5% cada uno, acorde a su aporte real.",
     "Debe existir aporte efectivo de B y C, aunque sea menor. "
     "La distribución debe ser económicamente justificada."],
    ["Alt. 4\nSpA régimen Pro Pyme (14D N°3)",
     "Constituir SpA con Pedro como único accionista. Tributación de la empresa al 25% "
     "de IDPC; Pedro paga IGC sobre retiros. Permite reinversión y planificación.",
     "No tiene restricciones de sociedad de profesionales. "
     "Requiere sustancia económica en la sociedad."],
    ["Alt. 5\nRégimen Transparente (14D N°8)",
     "Si los socios son solo personas naturales, la empresa no paga IDPC; "
     "cada socio paga IGC directo sobre su participación.",
     "ATENCIÓN: no soluciona el problema si B y C siguen sin aportar valor real. "
     "Solo válido con socios con sustancia."],
]
for i, f in enumerate(alternativas):
    row_data(t_alt, i + 1, f, bg="F4F6F9" if i % 2 == 0 else None)

doc.add_paragraph()


# ═══════════════════════════════════════════════════════════════════════════════
# VII. CONCLUSIONES
# ═══════════════════════════════════════════════════════════════════════════════
add_section("VII.  CONCLUSIONES")

conclusiones = [
    ("ABUSO, no simulación.",
     "El Caso 86 configura un abuso de las formas jurídicas (Art. 4 ter CT). "
     "La sociedad y los socios son reales, pero la forma es inapropiada: no hay "
     "sustancia económica en B y C. No hay acto secreto, por lo tanto no es simulación."),
    ("La sociedad D no califica.",
     "Por incumplimiento de los requisitos de las Circulares N°21/1991 y N°50/2020 "
     "(no todos los socios ejercen su profesión para la sociedad), D no es una "
     "sociedad de profesionales legítima bajo el Art. 42 N°2 LIR."),
    ("El SII puede actuar.",
     "Acreditados los elementos del Art. 4 ter CT, el Director puede requerir al TTA "
     "la declaración de abuso (Art. 4 quinquies CT). El efecto: atribución íntegra "
     "de los $300 MM a Pedro para su IGC, más intereses y multas."),
    ("Hay soluciones legítimas.",
     "Pedro puede organizar su actividad profesional a través de una sociedad, siempre "
     "que los socios tengan rol económico real y la distribución de utilidades refleje "
     "el aporte efectivo de cada uno."),
    ("El mensaje del caso es claro.",
     "El Derecho Tributario chileno acepta la planificación, pero exige que la forma "
     "jurídica tenga una justificación económica real más allá del ahorro fiscal. "
     "La NGA actúa precisamente cuando esa justificación no existe.", True),
]

for i, item in enumerate(conclusiones):
    titulo, texto = item[0], item[1]
    bold_final = len(item) > 2 and item[2]
    add_subtitle(f"{i+1}. {titulo}", color=AZUL if not bold_final else ROJO)
    add_para(texto, bold=bold_final)

doc.add_paragraph()

linea2 = doc.add_paragraph()
linea2.alignment = WD_ALIGN_PARAGRAPH.CENTER
linea2.add_run("─" * 80).font.color.rgb = CELESTE

pie = doc.add_paragraph()
pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
pie_r = pie.add_run(
    "Grupo 4: Andrea Añasco · Gema Sepúlveda · Karen Rebolledo · Nicolás Muñoz\n"
    "Magíster en Dirección Tributaria — UVM | Caso 86 · Catálogo SII 2025\n"
    "Fuentes: LIR Art. 42 N°2; CT Arts. 4 bis, 4 ter, 4 quáter, 4 quinquies; "
    "Circ. N°21/1991; Circ. N°50/2020; Circ. N°65/2015; Oficio N°2359/2022 SII."
)
pie_r.font.size   = Pt(8)
pie_r.font.italic = True
pie_r.font.color.rgb = GRIS

doc.save(OUTPUT)
print(f"✔ Informe Word generado: {OUTPUT}")
