#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Toma el PPT base de la compañera (8 slides) y completa las slides 7 y 8
con contenido de alternativas y conclusiones, respetando el estilo exacto.

Layout corregido — sin solapamientos:
  • El título en ambas slides ocupa top=0.70 h=1.50 → termina en y=2.20
  • Todo el contenido empieza desde y=2.22
  • Área izquierda : x=0.50  w=6.90  (hasta x=7.40)
  • Panel derecho  : x=7.50  w=5.60  (hasta x=13.10)
  • Sin <p:ph> conflictivos — solo auto_shapes y text_boxes normales
"""

from pptx import Presentation
from lxml import etree

INPUT  = "/home/user/Prueba/base_companera.pptx"
OUTPUT = "/home/user/Prueba/Caso86_Presentacion.pptx"

prs = Presentation(INPUT)

# ── Constantes de layout ──────────────────────────────────────────────────────
# El título termina en y ≈ 2.20 in — empezamos todo el contenido desde Y0
Y0   = 2.22   # inicio del área de contenido
Y_BOT = 7.28  # fondo disponible (margen de 0.22)

LX, LW = 0.50, 6.90   # área izquierda
PX, PW = 7.50, 5.60   # panel derecho
PH_total = Y_BOT - Y0  # 5.06 in de alto disponible

# Cabecera izquierda (dorada)
HDR_H  = 0.46
# 5 bloques (titulo+descripción)
BLK_TIT  = 0.32   # alto del rectángulo titulo de cada bloque
BLK_DESC = 0.52   # alto del rectángulo descripción
BLK_GAP  = 0.06   # espacio entre bloques
BLK_TOTAL = BLK_TIT + BLK_DESC + BLK_GAP  # 0.90 × 5 = 4.50

BOXES_START = Y0 + HDR_H + 0.04  # 2.22 + 0.46 + 0.04 = 2.72

def I(inches):
    """Inches → EMU."""
    return int(inches * 914400)

def _x(s):
    """Escape XML."""
    return (s.replace('&','&amp;').replace('<','&lt;')
             .replace('>','&gt;').replace('"','&quot;'))

def sp_rect_accent1(sid, x, y, w, h, lines, sz=1100, bold_first=True, anchor="ctr"):
    """Rectángulo con fill accent1 (azul del tema) y texto blanco — igual slide 6."""
    paras = ""
    for i, line in enumerate(lines):
        b = "<a:b>1</a:b>" if (i == 0 and bold_first) else ""
        paras += f"""
    <a:p><a:r><a:rPr lang="es-CL" sz="{sz}" dirty="0">{b}</a:rPr>
    <a:t>{_x(line)}</a:t></a:r></a:p>"""
    return etree.fromstring(f"""<p:sp
        xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:nvSpPr>
    <p:cNvPr id="{sid}" name="Rect{sid}"/>
    <p:cNvSpPr/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{I(x)}" y="{I(y)}"/>
    <a:ext cx="{I(w)}" cy="{I(h)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
  <p:style>
    <a:lnRef idx="2"><a:schemeClr val="accent1"><a:shade val="15000"/></a:schemeClr></a:lnRef>
    <a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef>
    <a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef>
    <a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef>
  </p:style>
  <p:txBody>
    <a:bodyPr rtlCol="0" anchor="{anchor}">
      <a:normAutofit/>
    </a:bodyPr>
    <a:lstStyle/>
    {paras}
  </p:txBody>
</p:sp>""")

def sp_rect_tint(sid, x, y, w, h, lines, sz=1000, anchor="ctr"):
    """Rectángulo con relleno accent1 tint 20% (azul claro) y texto oscuro — descripción."""
    paras = ""
    for line in lines:
        paras += f"""
    <a:p><a:pPr marL="114300" indent="-114300"/>
    <a:r><a:rPr lang="es-CL" sz="{sz}" dirty="0"/>
    <a:t>{_x(line)}</a:t></a:r></a:p>"""
    return etree.fromstring(f"""<p:sp
        xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:nvSpPr>
    <p:cNvPr id="{sid}" name="Desc{sid}"/>
    <p:cNvSpPr/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{I(x)}" y="{I(y)}"/>
    <a:ext cx="{I(w)}" cy="{I(h)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill>
      <a:schemeClr val="accent1"><a:tint val="20000"/></a:schemeClr>
    </a:solidFill>
  </p:spPr>
  <p:txBody>
    <a:bodyPr rtlCol="0" anchor="{anchor}">
      <a:normAutofit/>
    </a:bodyPr>
    <a:lstStyle/>
    {paras}
  </p:txBody>
</p:sp>""")

def sp_dark_panel(sid, x, y, w, h):
    """Panel oscuro degradado — igual Rectángulo 11 de slide 6."""
    return etree.fromstring(f"""<p:sp
        xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:nvSpPr>
    <p:cNvPr id="{sid}" name="Panel{sid}"/>
    <p:cNvSpPr/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{I(x)}" y="{I(y)}"/>
    <a:ext cx="{I(w)}" cy="{I(h)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:gradFill>
      <a:gsLst>
        <a:gs pos="0"><a:schemeClr val="dk1">
          <a:tint val="55000"/><a:satMod val="105000"/>
        </a:schemeClr></a:gs>
        <a:gs pos="100000"><a:schemeClr val="dk1">
          <a:tint val="70000"/><a:satMod val="100000"/>
        </a:schemeClr></a:gs>
      </a:gsLst>
      <a:lin ang="5400000" scaled="0"/>
    </a:gradFill>
  </p:spPr>
  <p:style>
    <a:lnRef idx="1"><a:schemeClr val="dk1"/></a:lnRef>
    <a:fillRef idx="2"><a:schemeClr val="dk1"/></a:fillRef>
    <a:effectRef idx="1"><a:schemeClr val="dk1"/></a:effectRef>
    <a:fontRef idx="minor"><a:schemeClr val="dk1"/></a:fontRef>
  </p:style>
  <p:txBody>
    <a:bodyPr/><a:lstStyle/>
    <a:p><a:endParaRPr lang="es-CL" dirty="0"/></a:p>
  </p:txBody>
</p:sp>""")

def sp_header_dark(sid, x, y, w, h, text):
    """Header negro/tx1 semitransparente con texto dorado — como 'Normativa Base' slide 4."""
    return etree.fromstring(f"""<p:sp
        xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:nvSpPr>
    <p:cNvPr id="{sid}" name="Hdr{sid}"/>
    <p:cNvSpPr/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{I(x)}" y="{I(y)}"/>
    <a:ext cx="{I(w)}" cy="{I(h)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:solidFill>
      <a:schemeClr val="tx1"><a:alpha val="82000"/></a:schemeClr>
    </a:solidFill>
  </p:spPr>
  <p:txBody>
    <a:bodyPr rtlCol="0" anchor="ctr"/>
    <a:lstStyle/>
    <a:p>
      <a:r>
        <a:rPr lang="es-CL" sz="1600" b="1" dirty="0">
          <a:solidFill><a:srgbClr val="FFC000"/></a:solidFill>
        </a:rPr>
        <a:t>{_x(text)}</a:t>
      </a:r>
    </a:p>
  </p:txBody>
</p:sp>""")

def tb_panel(sid, x, y, w, h, title, lines, sz=1000):
    """Sección dentro del panel oscuro: título dorado + líneas en blanco."""
    # título
    title_xml = f"""
    <a:p><a:r>
      <a:rPr lang="es-CL" sz="1200" b="1" dirty="0">
        <a:solidFill><a:srgbClr val="FFC000"/></a:solidFill>
      </a:rPr>
      <a:t>{_x(title)}</a:t>
    </a:r></a:p>
    <a:p><a:endParaRPr lang="es-CL" sz="600" dirty="0"/></a:p>"""
    # líneas de contenido
    for line in lines:
        title_xml += f"""
    <a:p><a:r>
      <a:rPr lang="es-CL" sz="{sz}" dirty="0">
        <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
      </a:rPr>
      <a:t>{_x(line)}</a:t>
    </a:r></a:p>"""
    return etree.fromstring(f"""<p:sp
        xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:nvSpPr>
    <p:cNvPr id="{sid}" name="TB{sid}"/>
    <p:cNvSpPr txBox="1"/><p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{I(x)}" y="{I(y)}"/>
    <a:ext cx="{I(w)}" cy="{I(h)}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:noFill/>
  </p:spPr>
  <p:txBody>
    <a:bodyPr wrap="square" rtlCol="0">
      <a:normAutofit/>
    </a:bodyPr>
    <a:lstStyle/>
    {title_xml}
  </p:txBody>
</p:sp>""")

def get_spTree(slide):
    P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
    cSld = slide._element.find(f'{{{P}}}cSld')
    return cSld.find(f'{{{P}}}spTree')


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Alternativas de solución
# ═══════════════════════════════════════════════════════════════════════════════
s7 = prs.slides[6]
t7 = get_spTree(s7)

# 1. Panel oscuro derecho — del Y0 al fondo
t7.append(sp_dark_panel(100, PX, Y0, PW, PH_total))

# 2. Contenido del panel derecho (dos secciones)
PAD = 0.14
t7.append(tb_panel(101, PX+PAD, Y0+0.12, PW-PAD*2, 2.35,
    "Consideraciones",
    [
        "Para que cualquier alternativa sea",
        "válida frente a las NGA, el socio",
        "debe poder acreditar SUSTANCIA",
        "ECONÓMICA REAL:",
        "",
        "• Todos los socios deben ejercer su",
        "  profesión para la sociedad.",
        "• La distribución de utilidades debe",
        "  reflejar el aporte efectivo.",
        "• El capital debe ser significativo",
        "  respecto a los ingresos.",
    ], sz=980))

t7.append(tb_panel(102, PX+PAD, Y0+2.60, PW-PAD*2, 2.35,
    "Normativa de referencia",
    [
        "• Art. 42 N°2 LIR",
        "  Rentas 2ª categoría",
        "• Art. 14 D N°3 y N°8 LIR",
        "  Regímenes Pro Pyme",
        "• Arts. 4 bis y 4 ter CT",
        "  Normas generales antielusivas",
        "• Circ. N°21/1991 y N°50/2020 SII",
        "• Circ. N°65/2015 SII",
    ], sz=980))

# 3. Header izquierdo (negro con texto dorado)
t7.append(sp_header_dark(103, LX, Y0, LW, HDR_H, "Alternativas de solución"))

# 4. Las 5 alternativas
alts = [
    ("Alt. 1 — Profesional independiente (Art. 42 N°2 LIR)",
     "A ejerce como persona natural, emite boletas de honorarios. Puede deducir gastos "
     "efectivos o presuntos (30%, tope 15 UTA). Paga IGC sobre el total. Sin riesgo NGA."),
    ("Alt. 2 — Sociedad de profesionales con socios que trabajen",
     "A se asocia con profesionales que efectivamente presten servicios para la sociedad. "
     "D califica legítimamente bajo Art. 42 N°2 inc. 3° LIR."),
    ("Alt. 3 — Distribución proporcional al trabajo real",
     "Estatutos con participación diferenciada (ej. A 90%, B 5%, C 5%) o sueldo patronal "
     "previo al reparto. La distribución debe reflejar el aporte efectivo de cada socio."),
    ("Alt. 4 — SpA — régimen Pro Pyme General (Art. 14 D N°3 LIR)",
     "SpA con IDPC al 25%. Permite planificación de retiros y reinversión. No tiene las "
     "restricciones de la sociedad de profesionales, pero exige sustancia económica real."),
    ("Alt. 5 — Régimen Pro Pyme Transparente (Art. 14 D N°8 LIR)",
     "La empresa no paga IDPC; socios tributan con IGC directo. Válido solo si B y C "
     "aportan sustancia real — de lo contrario el SII igualmente puede aplicar las NGA."),
]

for i, (titulo, desc) in enumerate(alts):
    y_blk = BOXES_START + i * BLK_TOTAL
    # Rectángulo título — accent1 azul, texto blanco
    t7.append(sp_rect_accent1(110 + i*2, LX, y_blk, LW, BLK_TIT,
                               [titulo], sz=1100, bold_first=True))
    # Rectángulo descripción — accent1 claro, texto oscuro
    t7.append(sp_rect_tint(111 + i*2, LX, y_blk + BLK_TIT, LW, BLK_DESC,
                            [desc], sz=1000))


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Conclusiones
# ═══════════════════════════════════════════════════════════════════════════════
s8 = prs.slides[7]
t8 = get_spTree(s8)

# 1. Panel oscuro derecho
t8.append(sp_dark_panel(200, PX, Y0, PW, PH_total))

# 2. Contenido del panel derecho
t8.append(tb_panel(201, PX+PAD, Y0+0.12, PW-PAD*2, 2.80,
    "Posición del Grupo 4",
    [
        "El Caso 86 configura ABUSO de",
        "formas jurídicas (Art. 4 ter CT).",
        "",
        "NO es simulación: los actos son",
        "reales, pero la forma jurídica es",
        "inapropiada — el único efecto",
        "relevante es tributario.",
        "",
        "Integrantes:",
        "  Andrea Añasco",
        "  Gema Sepúlveda",
        "  Karen Rebolledo",
        "  Nicolás Muñoz",
    ], sz=980))

t8.append(tb_panel(202, PX+PAD, Y0+3.05, PW-PAD*2, 1.90,
    "Norma aplicada",
    [
        "• Art. 4 ter CT — Abuso",
        "• Art. 4 quinquies CT — Procedimiento",
        "• Art. 42 N°2 LIR",
        "• Circ. N°21/1991 y N°50/2020 SII",
        "• Ley N°20.780 y Ley N°21.210",
    ], sz=980))

# 3. Header izquierdo
t8.append(sp_header_dark(203, LX, Y0, LW, HDR_H, "Síntesis del análisis — Caso 86"))

# 4. Las 5 conclusiones
concls = [
    ("1. ABUSO de formas jurídicas — Art. 4 ter CT",
     "El esquema usa la figura de sociedad de profesionales de forma inapropiada. "
     "Los actos son reales, pero el único efecto relevante distinto del tributario es ninguno."),
    ("2. La sociedad D no califica como sociedad de profesionales",
     "B y C no ejercen su profesión para D, incumpliendo el requisito copulativo "
     "de las Circulares N°21/1991 y N°50/2020 del SII."),
    ("3. El SII puede requerir declaración de abuso al TTA",
     "Art. 4 quinquies CT: el Director del SII puede pedir al Tribunal Tributario y "
     "Aduanero la declaración de abuso → recalificación + intereses + multas."),
    ("4. Existen alternativas legítimas de organización",
     "Profesional independiente, sociedad de profesionales con socios reales, "
     "SpA Pro Pyme (14 D N°3) o distribución de utilidades proporcional al trabajo."),
    ("5. La clave: sustancia económica real",
     "Toda estructura debe demostrar que los socios aportan trabajo, capital, clientela "
     "o dirección efectiva. La forma jurídica por sí sola no es suficiente frente a NGA."),
]

for i, (titulo, desc) in enumerate(concls):
    y_blk = BOXES_START + i * BLK_TOTAL
    t8.append(sp_rect_accent1(210 + i*2, LX, y_blk, LW, BLK_TIT,
                               [titulo], sz=1100, bold_first=True))
    t8.append(sp_rect_tint(211 + i*2, LX, y_blk + BLK_TIT, LW, BLK_DESC,
                            [desc], sz=1000))


prs.save(OUTPUT)
print(f"✔  PPT guardado: {OUTPUT}")

# Verificar que no hay solapamientos en slides 7 y 8
print("\nVerificación de posiciones (top + height):")
prs2 = Presentation(OUTPUT)
for si in [6, 7]:
    sl = prs2.slides[si]
    print(f"\n  Slide {si+1}:")
    shapes = [(sh.top/914400, sh.height/914400, sh.name) for sh in sl.shapes]
    for t, h, n in sorted(shapes):
        print(f"    y={t:5.2f}  bot={t+h:5.2f}  h={h:5.2f}  {n}")
